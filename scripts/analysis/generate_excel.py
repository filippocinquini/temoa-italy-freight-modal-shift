import sqlite3
import re
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# File paths
db_path = 'TRA_MS_OFF_BAU.sqlite'
excel_path = 'TRA_ROA_HTR_MTR_LCV_costs.xlsx'

# Regex pattern for matching technologies
tech_pattern = re.compile(r'^TRA_ROA_(HTR|MTR|LCV)_([A-Z0-9]+)_N_D([1-5])$')

# Mappings for human readability
vehicle_map = {
    'HTR': 'Heavy Trucks (HTR)',
    'MTR': 'Medium Trucks (MTR)',
    'LCV': 'Light Commercial Vehicles (LCV)'
}

powertrain_map = {
    'DST': 'Diesel',
    'ELC': 'Full Electric',
    'FCELL': 'Fuel Cell',
    'FULHYB': 'Full Hybrid',
    'PHEV': 'Plug-in Hybrid',
    'LPG': 'LPG',
    'NGA': 'Natural Gas',
    'GSL': 'Gasoline',
    'BIO': 'Biofuel'
}

def parse_tech(tech_id):
    m = tech_pattern.match(tech_id)
    if m:
        veh_code = m.group(1)
        pt_code = m.group(2)
        dist_code = f"D{m.group(3)}"
        
        veh_label = vehicle_map.get(veh_code, veh_code)
        pt_label = powertrain_map.get(pt_code, pt_code)
        
        return veh_label, pt_label, dist_code
    return "Unknown", "Unknown", "Unknown"

# 1. Connect to SQLite and fetch data
print("Connecting to database...")
conn = sqlite3.connect(db_path)

# CostInvest
print("Fetching CostInvest...")
query_invest = """
    SELECT c.regions, c.tech, t.tech_desc, c.vintage, c.cost_invest, c.cost_invest_units, c.cost_invest_notes
    FROM CostInvest c
    JOIN technologies t ON c.tech = t.tech
"""
df_invest = pd.read_sql_query(query_invest, conn)

# CostFixed
print("Fetching CostFixed...")
query_fixed = """
    SELECT c.regions, c.periods, c.tech, t.tech_desc, c.vintage, c.cost_fixed, c.cost_fixed_units, c.cost_fixed_notes
    FROM CostFixed c
    JOIN technologies t ON c.tech = t.tech
"""
df_fixed = pd.read_sql_query(query_fixed, conn)

# CostVariable
print("Fetching CostVariable...")
query_variable = """
    SELECT c.regions, c.periods, c.tech, t.tech_desc, c.vintage, c.cost_variable, c.cost_variable_units, c.cost_variable_notes
    FROM CostVariable c
    JOIN technologies t ON c.tech = t.tech
"""
df_variable = pd.read_sql_query(query_variable, conn)

conn.close()

# 2. Filter data
print("Filtering target technologies...")
df_invest_filtered = df_invest[df_invest['tech'].apply(lambda x: bool(tech_pattern.match(x)))].copy()
df_fixed_filtered = df_fixed[df_fixed['tech'].apply(lambda x: bool(tech_pattern.match(x)))].copy()
df_variable_filtered = df_variable[df_variable['tech'].apply(lambda x: bool(tech_pattern.match(x)))].copy()

# Add parsed columns
for df in [df_invest_filtered, df_fixed_filtered, df_variable_filtered]:
    parsed = df['tech'].apply(parse_tech)
    df['Vehicle Type'] = [p[0] for p in parsed]
    df['Powertrain'] = [p[1] for p in parsed]
    df['Distance Step'] = [p[2] for p in parsed]

# Reorder and rename columns
df_invest_filtered = df_invest_filtered[[
    'regions', 'tech', 'tech_desc', 'Vehicle Type', 'Powertrain', 'Distance Step', 
    'vintage', 'cost_invest', 'cost_invest_units', 'cost_invest_notes'
]]
df_invest_filtered.rename(columns={
    'regions': 'Region',
    'tech': 'Tech ID',
    'tech_desc': 'Description',
    'vintage': 'Vintage Year',
    'cost_invest': 'Investment Cost',
    'cost_invest_units': 'Units',
    'cost_invest_notes': 'Notes'
}, inplace=True)

df_fixed_filtered = df_fixed_filtered[[
    'regions', 'tech', 'tech_desc', 'Vehicle Type', 'Powertrain', 'Distance Step',
    'periods', 'vintage', 'cost_fixed', 'cost_fixed_units', 'cost_fixed_notes'
]]
df_fixed_filtered.rename(columns={
    'regions': 'Region',
    'tech': 'Tech ID',
    'tech_desc': 'Description',
    'periods': 'Period Year',
    'vintage': 'Vintage Year',
    'cost_fixed': 'Fixed Cost',
    'cost_fixed_units': 'Units',
    'cost_fixed_notes': 'Notes'
}, inplace=True)

df_variable_filtered = df_variable_filtered[[
    'regions', 'tech', 'tech_desc', 'Vehicle Type', 'Powertrain', 'Distance Step',
    'periods', 'vintage', 'cost_variable', 'cost_variable_units', 'cost_variable_notes'
]]
df_variable_filtered.rename(columns={
    'regions': 'Region',
    'tech': 'Tech ID',
    'tech_desc': 'Description',
    'periods': 'Period Year',
    'vintage': 'Vintage Year',
    'cost_variable': 'Variable Cost',
    'cost_variable_units': 'Units',
    'cost_variable_notes': 'Notes'
}, inplace=True)

# 3. Build Excel Workbook with premium styling
print("Generating styled Excel workbook...")
wb = Workbook()

# Set up sheets
ws_overview = wb.active
ws_overview.title = "Overview"
ws_invest = wb.create_sheet(title="CostInvest")
ws_fixed = wb.create_sheet(title="CostFixed")
ws_variable = wb.create_sheet(title="CostVariable")

# Colors
NAVY_HEADER = "1F4E78"
ICE_BLUE = "DDEBF7"
LIGHT_GRAY_ZEBRA = "F9FBFD"
WHITE = "FFFFFF"
DARK_GRAY_TEXT = "333333"

# Fonts
font_title = Font(name="Segoe UI", size=18, bold=True, color="1F4E78")
font_subtitle = Font(name="Segoe UI", size=12, italic=True, color="595959")
font_section = Font(name="Segoe UI", size=14, bold=True, color="1F4E78")
font_header = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
font_body = Font(name="Segoe UI", size=11, color="000000")
font_body_bold = Font(name="Segoe UI", size=11, bold=True, color="000000")
font_small = Font(name="Segoe UI", size=9, italic=True, color="595959")

# Fills
fill_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
fill_accent = PatternFill(start_color=ICE_BLUE, end_color=ICE_BLUE, fill_type="solid")
fill_zebra = PatternFill(start_color=LIGHT_GRAY_ZEBRA, end_color=LIGHT_GRAY_ZEBRA, fill_type="solid")
fill_white = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")

# Alignments
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

# Borders
thin_side = Side(border_style="thin", color="D9D9D9")
border_thin = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
border_header = Border(left=thin_side, right=thin_side, top=thin_side, bottom=Side(border_style="medium", color="1F4E78"))

# Write Overview sheet
print("Writing Overview sheet...")
ws_overview.views.sheetView[0].showGridLines = True

ws_overview.column_dimensions['B'].width = 25
ws_overview.column_dimensions['C'].width = 45

# Title Block
ws_overview['B2'] = "TEMOA COST EXTRACTION REPORT"
ws_overview['B2'].font = font_title
ws_overview['B3'] = f"Extracted on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
ws_overview['B3'].font = font_subtitle

# Metadata
ws_overview['B5'] = "Metadata"
ws_overview['B5'].font = font_section
metadata = [
    ("Source Database", "TRA_MS_OFF_BAU.sqlite"),
    ("Target Technologies", "TRA_ROA_HTR/MTR/LCV (New, D1-D5)"),
    ("Pattern Matched", "TRA_ROA_(HTR|MTR|LCV)_*_N_D[1-5]"),
    ("Total Unique Techs", len(df_invest_filtered['Tech ID'].unique())),
]
for idx, (k, v) in enumerate(metadata):
    row = 6 + idx
    ws_overview.cell(row=row, column=2, value=k).font = font_body_bold
    ws_overview.cell(row=row, column=3, value=v).font = font_body
    ws_overview.cell(row=row, column=2).border = border_thin
    ws_overview.cell(row=row, column=3).border = border_thin

# Summary stats
ws_overview['B12'] = "Extracted Cost Entries"
ws_overview['B12'].font = font_section

headers_summary = ["Cost Category", "Extracted Rows", "Unique Technologies"]
for col_idx, h in enumerate(headers_summary):
    cell = ws_overview.cell(row=13, column=2 + col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = border_header

summary_data = [
    ("Investment Cost (CostInvest)", len(df_invest_filtered), df_invest_filtered['Tech ID'].nunique()),
    ("Fixed O&M Cost (CostFixed)", len(df_fixed_filtered), df_fixed_filtered['Tech ID'].nunique()),
    ("Variable Cost (CostVariable)", len(df_variable_filtered), df_variable_filtered['Tech ID'].nunique())
]

for row_offset, (cat, rows, techs) in enumerate(summary_data):
    row = 14 + row_offset
    c1 = ws_overview.cell(row=row, column=2, value=cat)
    c2 = ws_overview.cell(row=row, column=3, value=rows)
    c3 = ws_overview.cell(row=row, column=4, value=techs)
    
    c1.font = font_body_bold
    c2.font = font_body
    c3.font = font_body
    
    c1.border = border_thin
    c2.border = border_thin
    c3.border = border_thin
    
    c2.alignment = align_right
    c3.alignment = align_right

# Abbreviations Legend
ws_overview['B19'] = "Legend & Abbreviations"
ws_overview['B19'].font = font_section

legend_data = [
    ("HTR", "Heavy Trucks"),
    ("MTR", "Medium Trucks"),
    ("LCV", "Light Commercial Vehicles"),
    ("DST", "Diesel"),
    ("ELC", "Full Electric"),
    ("FCELL", "Hydrogen Fuel Cell"),
    ("FULHYB", "Gasoline/Diesel Full Hybrid"),
    ("PHEV", "Plug-in Hybrid"),
    ("LPG", "Liquefied Petroleum Gas"),
    ("NGA", "Natural Gas"),
    ("GSL", "Gasoline"),
    ("BIO", "Biofuel"),
    ("D1 - D5", "Distance bands (e.g. D1: <50 km, D2: 50-100 km, D3: 100-300 km, D4: 300-500 km, D5: >500 km)"),
    ("N", "New Technology")
]

ws_overview.cell(row=20, column=2, value="Code").font = font_header
ws_overview.cell(row=20, column=2).fill = fill_header
ws_overview.cell(row=20, column=2).border = border_header
ws_overview.cell(row=20, column=2).alignment = align_center

ws_overview.cell(row=20, column=3, value="Description / Mapping").font = font_header
ws_overview.cell(row=20, column=3).fill = fill_header
ws_overview.cell(row=20, column=3).border = border_header
ws_overview.cell(row=20, column=3).alignment = align_center

for row_offset, (code, desc) in enumerate(legend_data):
    row = 21 + row_offset
    c1 = ws_overview.cell(row=row, column=2, value=code)
    c2 = ws_overview.cell(row=row, column=3, value=desc)
    
    c1.font = font_body_bold
    c2.font = font_body
    
    c1.border = border_thin
    c2.border = border_thin
    c1.alignment = align_center

# Helper function to write cost sheets
def write_cost_sheet(ws, df, numeric_col_name):
    ws.views.sheetView[0].showGridLines = True
    # Freeze top row
    ws.freeze_panes = "A2"
    
    # Write headers
    headers = list(df.columns)
    for col_idx, h in enumerate(headers):
        cell = ws.cell(row=1, column=1 + col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header
    
    # Write data
    for row_idx, row_values in enumerate(df.values):
        row_num = 2 + row_idx
        fill = fill_zebra if row_idx % 2 == 1 else fill_white
        for col_idx, val in enumerate(row_values):
            cell = ws.cell(row=row_num, column=1 + col_idx, value=val)
            cell.font = font_body
            cell.fill = fill
            cell.border = border_thin
            
            # Alignments & formats
            col_name = headers[col_idx]
            if col_name in ['Region', 'Distance Step', 'Vintage Year', 'Period Year']:
                cell.alignment = align_center
            elif col_name == numeric_col_name:
                cell.alignment = align_right
                cell.number_format = '#,##0.00'
            else:
                cell.alignment = align_left
                
    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Check length of string
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Apply width
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

print("Writing CostInvest sheet...")
write_cost_sheet(ws_invest, df_invest_filtered, 'Investment Cost')

print("Writing CostFixed sheet...")
write_cost_sheet(ws_fixed, df_fixed_filtered, 'Fixed Cost')

print("Writing CostVariable sheet...")
write_cost_sheet(ws_variable, df_variable_filtered, 'Variable Cost')

# Save workbook
print(f"Saving workbook to {excel_path}...")
wb.save(excel_path)
print("Excel workbook created successfully!")
